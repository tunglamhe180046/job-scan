"""
Create Word Document (.docx) for Ly Lich Trich Ngang - Nguyen Tung Lam
Uses python-docx to build a perfectly styled, 100% pixel-perfect aligned Word Document.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    """Set background color for a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal padding for table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    """Set sharp black borders for all cells in a table."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), val)
        node.set(qn('w:sz'), sz)
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), color)
        tblBorders.append(node)
    tblPr.append(tblBorders)

def create_resume_docx(output_path=None):
    if output_path is None:
        # Default to root directory single docx file
        base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
        output_path = os.path.join(base_dir, "Ly_Lich_Trich_Ngang_Nguyen_Tung_Lam.docx")
        
    doc = Document()
    
    # Page Margins: 0.75 inch (54 pt)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
    # --- HEADER SECTION (Full Width Centered Header + Horizontal Meta Line) ---
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(7.0)
    
    cell_title = header_table.cell(0, 0)
    
    # Line 1: Company Name
    p_comp = cell_title.paragraphs[0]
    p_comp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_comp.paragraph_format.space_after = Pt(2)
    run_comp = p_comp.add_run("CÔNG TY CỔ PHẦN TẬP ĐOÀN MINH BẢO")
    run_comp.font.name = "Arial"
    run_comp.font.size = Pt(13)
    run_comp.font.bold = True
    run_comp.font.color.rgb = RGBColor(26, 54, 93) # Dark Navy
    
    # Line 2: Document Title
    p_title = cell_title.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("LÝ LỊCH TRÍCH NGANG")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(197, 48, 48) # Crimson Red

    # Line 3: Horizontal Meta Info (Mã, Ngày BH, Sửa đổi on 1 SINGLE line)
    p_meta = cell_title.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(2)
    run_meta = p_meta.add_run("Mã: ....................        Ngày BH: ....................        Sửa đổi: ....................")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(113, 128, 150)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SECTION HELPER ---
    def add_section_header(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title_text)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(43, 108, 176)

    # --- SECTION 1: THÔNG TIN CÁ NHÂN ---
    add_section_header("I. THÔNG TIN CÁ NHÂN (*)")
    
    personal_data = [
        ("Họ và tên đầy đủ:", "Nguyễn Tùng Lâm"),
        ("Ngày sinh:", "12/08/2004"),
        ("Giới tính / Hôn nhân:", "Nam  |  Độc thân"),
        ("Nguyên quán / Nơi sinh:", "21 Lê Đức Thọ, Mỹ Đình 2, Nam Từ Liêm, Hà Nội"),
        ("Số CCCD / CMND:", "001204027449  (Ngày cấp: 17/05/2021 tại Cục CSQLHC về TTXH)"),
        ("Hộ khẩu thường trú:", "Tổ dân phố số 10, Mỹ Đình 2, Nam Từ Liêm, Hà Nội"),
        ("Nơi ở hiện nay:", "21 Lê Đức Thọ, Mỹ Đình 2, Nam Từ Liêm, Hà Nội"),
        ("Số điện thoại liên hệ:", "0985469702"),
        ("Địa chỉ Email:", "tunglam07678@mail.com"),
        ("Mã số thuế cá nhân:", "8889943189"),
        ("Tài khoản ngân hàng:", "1065997321 - TECHCOMBANK"),
        ("Kỹ năng nổi bật (AI Agent):", "• Thành thạo các công cụ AI (Claude, Codex, Cursor, Antigravity Pro).\n• Khả năng xây dựng Skill, Workflow & MCP Server kết nối trực tiếp với ứng dụng (Unity, Web, API).")
    ]
    
    t1 = doc.add_table(rows=len(personal_data), cols=2)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.columns[0].width = Inches(2.2)
    t1.columns[1].width = Inches(4.8)
    
    for i, (label, val) in enumerate(personal_data):
        row = t1.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        
        # Label
        set_cell_background(c0, "F7FAFC")
        set_cell_margins(c0, top=100, bottom=100, left=120, right=120)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.bold = True
        
        # Value
        set_cell_margins(c1, top=100, bottom=100, left=120, right=120)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        
    set_table_borders(t1)

    # --- SECTION 2: TRÌNH ĐỘ HỌC VẤN ---
    add_section_header("II. TRÌNH ĐỘ HỌC VẤN VÀ ĐÀO TẠO")
    
    edu_data = [
        ("Trình độ học vấn:", "Đại học (Chính quy)"),
        ("Trường / Đơn vị đào tạo:", "Đại học FPT - Chuyên ngành Công nghệ thông tin (Kỹ sư cầu nối)"),
        ("Xếp loại tốt nghiệp:", "Khá (Tốt nghiệp năm 2026)"),
        ("Trình độ Ngoại ngữ:", "Tiếng Anh: IELTS 5.5  |  Tiếng Nhật: JLPT N4")
    ]
    
    t2 = doc.add_table(rows=len(edu_data), cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.columns[0].width = Inches(2.2)
    t2.columns[1].width = Inches(4.8)
    
    for i, (label, val) in enumerate(edu_data):
        row = t2.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_background(c0, "F7FAFC")
        set_cell_margins(c0, top=100, bottom=100, left=120, right=120)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(10)
        r0.font.bold = True
        
        set_cell_margins(c1, top=100, bottom=100, left=120, right=120)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        
    set_table_borders(t2)

    # --- SECTION 3: KINH NGHIỆM LÀM VIỆC ---
    add_section_header("III. KINH NGHIỆM LÀM VIỆC")
    
    exp_headers = ["Thời gian", "Tên Công ty", "Chức danh", "Mô tả công việc & Thành tựu nổi bật"]
    t3 = doc.add_table(rows=3, cols=4)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t3.columns[0].width = Inches(1.3)
    t3.columns[1].width = Inches(1.4)
    t3.columns[2].width = Inches(1.4)
    t3.columns[3].width = Inches(2.9)
    
    # Header Row
    hdr_cells = t3.rows[0].cells
    for j, h_text in enumerate(exp_headers):
        set_cell_background(hdr_cells[j], "2B6CB0")
        set_cell_margins(hdr_cells[j], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[j].paragraphs[0]
        if j < 3:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    exp_rows = [
        ("09/2025 -\n12/2025", "NIC Group", ".NET Developer", "• Phát triển WinForms (C#) và Web JS cho khách hàng Nhật Bản.\n• Bảo trì và nâng cấp các hệ thống cũ."),
        ("11/02/2026 -\n31/07/2026", "KAOPIZ Software", "Frontend Developer", "• Phát triển UI ứng dụng với Unity 2D 2022 (UGUI), tích hợp API.\n• Tự thiết kế MCP Server riêng kết nối Unity 2022 với Claude AI.\n• Giúp tối ưu hóa và giảm 50% thời gian phát triển ứng dụng.")
    ]
    
    for i, row_data in enumerate(exp_rows, start=1):
        r_cells = t3.rows[i].cells
        for j, text_val in enumerate(row_data):
            set_cell_margins(r_cells[j], top=100, bottom=100, left=100, right=100)
            p = r_cells[j].paragraphs[0]
            if j < 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text_val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            
    set_table_borders(t3)

    # --- SECTION 4: MONG MUỐN VÀ CAM ĐOAN ---
    add_section_header("IV. MONG MUỐN VÀ CAM ĐOAN")
    
    p_desire = doc.add_paragraph()
    p_desire.paragraph_format.space_after = Pt(10)
    r_d = p_desire.add_run("• Vị trí mong muốn tại CÔNG TY CỔ PHẦN TẬP ĐOÀN MINH BẢO: Dev (Khối IT)\n"
                           "• Ngày bắt đầu làm việc dự kiến: 03/08/2026\n"
                           "• Mức lương khởi điểm: 10.000.000 VNĐ (Gross)  |  Chính thức: 12.000.000 VNĐ (Gross)\n"
                           "• Cam đoan: Thông tin khai báo hoàn toàn chính xác, không vướng mắc trách nhiệm pháp lý với công ty cũ.")
    r_d.font.name = "Arial"
    r_d.font.size = Pt(10)
    
    # Save Document
    doc.save(output_path)
    print(f"Document successfully created at: {os.path.abspath(output_path)}")
    return os.path.abspath(output_path)

if __name__ == "__main__":
    create_resume_docx()
